"""Word 文档渲染核心

把「结构化文档描述（dict）」渲染成 .docx。所有排版能力集中在这里，
供 docx_create / docx_edit / report_gen 三个工具共用。

结构化描述格式（JSON）
---------------------
{
  "template": "report",                    # 可选，叠加 data/templates/<name>.yaml 的默认排版
  "cover": {"title": "...", "subtitle": "...",
            "fields": [{"label": "项目", "value": "..."}], "date": "2026-01-01"},
  "toc": true,                             # 插入自动目录（Word 打开时更新域）
  "header": "页眉文字", "footer_note": "页脚左侧文字", "page_number": true,
  "sections": [
    {"type": "heading", "level": 1, "text": "第一章"},
    {"type": "paragraph", "text": "...", "align": "justify",
     "first_line_indent": true, "line_spacing": 1.5, "bold": false, "font_size": 12},
    {"type": "list", "ordered": false, "items": ["甲", "乙"]},
    {"type": "table", "header": ["列1","列2"], "rows": [["a","b"]],
     "merged": [{"row": 0, "cols": [0, 1]}], "header_repeat": true, "caption": "表1"},
    {"type": "image", "path": "a.png", "width_cm": 12, "caption": "图1"},
    {"type": "formula", "omml": "<m:oMath>...</m:oMath>"},
    {"type": "formula", "text": "GM = KM - KG"},
    {"type": "page_break"},
    {"type": "section_break", "orientation": "landscape"}
  ]
}

已确认不支持的项
----------------
- 脚注：python-docx 无原生 API，本项目明确不提供（需求方已决定取消）。
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from shared.utils.logging import get_logger

logger = get_logger(__name__)

# OMML 命名空间（公式）
_OMML_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"

ALIGN_MAP = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    "both": WD_ALIGN_PARAGRAPH.JUSTIFY,
}

# 默认排版（可被模板覆盖）
DEFAULT_STYLE: dict[str, Any] = {
    "page": {"margin_cm": {"top": 2.54, "bottom": 2.54, "left": 3.17, "right": 3.17}},
    "body": {"font": "宋体", "font_size": 12, "line_spacing": 1.5, "first_line_indent_chars": 2},
    "heading": {
        "1": {"font": "黑体", "font_size": 16, "space_before": 12, "space_after": 6},
        "2": {"font": "黑体", "font_size": 14, "space_before": 8, "space_after": 4},
        "3": {"font": "黑体", "font_size": 12.5, "space_before": 6, "space_after": 3},
        "4": {"font": "黑体", "font_size": 12, "space_before": 4, "space_after": 2},
    },
    "cover": {"title_size": 26, "subtitle_size": 16, "field_size": 14, "font": "黑体"},
}


def _set_east_asian(run, font_name: str) -> None:
    """设置中文字体（python-docx 需要单独设置 eastAsia）"""
    run.font.name = font_name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), font_name)
    rfonts.set(qn("w:ascii"), font_name)
    rfonts.set(qn("w:hAnsi"), font_name)


def _add_field(paragraph, instr: str, placeholder: str = "") -> None:
    """向段落插入一个 Word 域（用于页码 / 目录）"""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr_el = OxmlElement("w:instrText")
    instr_el.set(qn("xml:space"), "preserve")
    instr_el.text = instr
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    text_el = OxmlElement("w:t")
    text_el.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for el in (begin, instr_el, sep, text_el, end):
        run._r.append(el)


def _enable_update_fields(doc) -> None:
    """让 Word 打开文档时自动更新域（目录、页码）"""
    settings = doc.settings.element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)


def _apply_template(style: dict[str, Any], template_name: str | None) -> dict[str, Any]:
    """叠加模板默认值（data/templates/<name>.yaml）"""
    merged = {k: (v.copy() if isinstance(v, dict) else v) for k, v in DEFAULT_STYLE.items()}
    if not template_name:
        return merged
    try:
        import yaml

        from shared.utils.config import get_settings

        settings = get_settings()
        data_dir = (getattr(settings, "DATA_DIR", "") or "").strip()
        base = Path(data_dir) if data_dir else (settings.PROJECT_ROOT / "data")
        path = base / "templates" / f"{template_name}.yaml"
        if not path.exists():
            logger.warning("模板文件不存在，使用内置默认排版", path=str(path))
            return merged
        data = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
        for key, value in data.items():
            if isinstance(value, dict) and isinstance(merged.get(key), dict):
                merged[key].update(value)
            else:
                merged[key] = value
        logger.info("已加载排版模板", name=template_name, path=str(path))
    except Exception as exc:  # noqa: BLE001
        logger.warning("模板加载失败，使用内置默认排版", error=str(exc))
    return merged


def _setup_page(doc, style: dict[str, Any]) -> None:
    section = doc.sections[0]
    margin = (style.get("page") or {}).get("margin_cm") or {}
    for side, setter in (
        ("top", section.top_margin), ("bottom", section.bottom_margin),
        ("left", section.left_margin), ("right", section.right_margin),
    ):
        if side in margin:
            setattr(section, f"{side}_margin", Cm(float(margin[side])))


def _render_cover(doc, cover: dict, style: dict[str, Any]) -> None:
    cfg = style.get("cover") or {}
    for _ in range(4):
        doc.add_paragraph()
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run(str(cover.get("title", "")))
    run.font.size = Pt(float(cfg.get("title_size", 26)))
    run.bold = True
    _set_east_asian(run, cfg.get("font", "黑体"))

    subtitle = cover.get("subtitle")
    if subtitle:
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r2 = p2.add_run(str(subtitle))
        r2.font.size = Pt(float(cfg.get("subtitle_size", 16)))
        _set_east_asian(r2, cfg.get("font", "黑体"))

    for _ in range(4):
        doc.add_paragraph()

    for item in cover.get("fields") or []:
        p3 = doc.add_paragraph()
        p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r3 = p3.add_run(f"{item.get('label', '')}：{item.get('value', '')}")
        r3.font.size = Pt(float(cfg.get("field_size", 14)))
        _set_east_asian(r3, cfg.get("font", "黑体"))

    date_text = cover.get("date")
    if date_text:
        doc.add_paragraph()
        p4 = doc.add_paragraph()
        p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r4 = p4.add_run(str(date_text))
        r4.font.size = Pt(float(cfg.get("field_size", 14)))
        _set_east_asian(r4, cfg.get("font", "黑体"))

    doc.add_page_break()


def _render_toc(doc) -> None:
    heading = doc.add_paragraph()
    run = heading.add_run("目　录")
    run.bold = True
    run.font.size = Pt(16)
    _set_east_asian(run, "黑体")
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

    toc_para = doc.add_paragraph()
    _add_field(toc_para, 'TOC \\o "1-3" \\h \\z \\u', "（请在 Word 中右键选择「更新域」生成目录）")
    doc.add_page_break()


def _render_header_footer(doc, spec: dict) -> None:
    section = doc.sections[0]
    header_text = spec.get("header")
    if header_text:
        hp = section.header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        hr = hp.add_run(str(header_text))
        hr.font.size = Pt(9)
        _set_east_asian(hr, "宋体")

    footer_note = spec.get("footer_note")
    show_page = spec.get("page_number", True)
    fp = section.footer.paragraphs[0]
    if footer_note:
        fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fr = fp.add_run(str(footer_note) + "    ")
        fr.font.size = Pt(9)
        _set_east_asian(fr, "宋体")
    if show_page:
        if not footer_note:
            fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _add_field(fp, " PAGE ", "1")


def _render_heading(doc, block: dict, style: dict[str, Any]) -> None:
    level = max(1, min(4, int(block.get("level", 1))))
    cfg = (style.get("heading") or {}).get(str(level)) or {}
    heading = doc.add_heading(level=level)
    run = heading.add_run(str(block.get("text", "")))
    run.font.size = Pt(float(cfg.get("font_size", 14)))
    run.bold = True
    _set_east_asian(run, cfg.get("font", "黑体"))
    pf = heading.paragraph_format
    if "space_before" in cfg:
        pf.space_before = Pt(float(cfg["space_before"]))
    if "space_after" in cfg:
        pf.space_after = Pt(float(cfg["space_after"]))


def _render_paragraph(doc, block: dict, style: dict[str, Any]) -> None:
    body = style.get("body") or {}
    para = doc.add_paragraph()
    text = str(block.get("text", ""))
    run = para.add_run(text)

    size = float(block.get("font_size", body.get("font_size", 12)))
    run.font.size = Pt(size)
    run.bold = bool(block.get("bold", False))
    _set_east_asian(run, block.get("font", body.get("font", "宋体")))
    if block.get("italic"):
        run.italic = True
    color = block.get("color")
    if color:
        try:
            run.font.color.rgb = RGBColor.from_string(str(color).lstrip("#").upper())
        except Exception:  # noqa: BLE001
            pass

    para.alignment = ALIGN_MAP.get(str(block.get("align", "justify")).lower(), WD_ALIGN_PARAGRAPH.JUSTIFY)

    pf = para.paragraph_format
    pf.line_spacing = float(block.get("line_spacing", body.get("line_spacing", 1.5)))
    if block.get("space_before") is not None:
        pf.space_before = Pt(float(block["space_before"]))
    if block.get("space_after") is not None:
        pf.space_after = Pt(float(block["space_after"]))
    first_indent = block.get("first_line_indent", True)
    if first_indent:
        chars = float(block.get("first_line_indent_chars", body.get("first_line_indent_chars", 2)))
        pf.first_line_indent = Pt(size * chars)


def _render_list(doc, block: dict, style: dict[str, Any]) -> None:
    ordered = bool(block.get("ordered", False))
    list_style = "List Number" if ordered else "List Bullet"
    for item in block.get("items") or []:
        try:
            para = doc.add_paragraph(style=list_style)
        except KeyError:
            para = doc.add_paragraph()
        run = para.add_run(str(item))
        run.font.size = Pt(float((style.get("body") or {}).get("font_size", 12)))
        _set_east_asian(run, (style.get("body") or {}).get("font", "宋体"))


def _render_table(doc, block: dict, style: dict[str, Any]) -> None:
    header = list(block.get("header") or [])
    rows = list(block.get("rows") or [])
    cols = max([len(header)] + [len(r) for r in rows] + [1])
    table = doc.add_table(rows=len(rows) + (1 if header else 0), cols=cols)
    try:
        table.style = "Table Grid"
    except KeyError:
        pass

    body_cfg = style.get("body") or {}
    font = body_cfg.get("font", "宋体")
    size = float(block.get("font_size", body_cfg.get("font_size", 10.5)))

    def _fill(cell, text: str, bold: bool = False) -> None:
        cell.text = ""
        para = cell.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run(str(text))
        run.font.size = Pt(size)
        run.bold = bold
        _set_east_asian(run, font)

    offset = 0
    if header:
        for idx, text in enumerate(header):
            _fill(table.cell(0, idx), text, bold=True)
        offset = 1

    for r_idx, row in enumerate(rows):
        for c_idx in range(cols):
            value = row[c_idx] if c_idx < len(row) else ""
            _fill(table.cell(r_idx + offset, c_idx), value)

    # 合并单元格
    for item in block.get("merged") or []:
        try:
            r = int(item["row"]) + offset
            first, last = int(item["cols"][0]), int(item["cols"][-1])
            table.cell(r, first).merge(table.cell(r, last))
        except Exception as exc:  # noqa: BLE001
            logger.warning("合并单元格失败", error=str(exc))

    # 表头跨页重复
    if header and block.get("header_repeat", True):
        try:
            tr = table.rows[0]._tr
            trPr = tr.get_or_add_trPr()
            tblHeader = OxmlElement("w:tblHeader")
            tblHeader.set(qn("w:val"), "true")
            trPr.append(tblHeader)
        except Exception:  # noqa: BLE001
            pass

    caption = block.get("caption")
    if caption:
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cr = cap.add_run(str(caption))
        cr.font.size = Pt(10.5)
        _set_east_asian(cr, font)


def _render_image(doc, block: dict, base_dir: Path | None) -> None:
    raw = str(block.get("path") or "")
    path = Path(raw)
    if not path.is_absolute() and base_dir is not None:
        path = base_dir / raw
    if not path.exists():
        raise FileNotFoundError(f"图片不存在：{raw}")
    width = Cm(float(block.get("width_cm", 12)))
    doc.add_picture(str(path), width=width)
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption = block.get("caption")
    if caption:
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cr = cap.add_run(str(caption))
        cr.font.size = Pt(10.5)
        _set_east_asian(cr, "宋体")


def _render_formula(doc, block: dict) -> None:
    omml = block.get("omml")
    if omml:
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        try:
            para._p.append(parse_xml(str(omml)))
        except Exception as exc:  # noqa: BLE001
            raise ValueError(f"OMML 公式解析失败：{exc}") from exc
        return
    text = block.get("text")
    if text:
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run(str(text))
        run.italic = True
        _set_east_asian(run, "Cambria Math")
        return
    raise ValueError("formula 需要 omml 或 text 其一")


def _render_section_break(doc, block: dict) -> None:
    """分节；orientation=landscape 时新节横向"""
    from docx.enum.section import WD_SECTION

    new_section = doc.add_section(WD_SECTION.NEW_PAGE)
    orientation = str(block.get("orientation", "portrait")).lower()
    if orientation == "landscape":
        width, height = new_section.page_width, new_section.page_height
        new_section.orientation = WD_ORIENT.LANDSCAPE
        new_section.page_width = height
        new_section.page_height = width


def render(spec: dict, base_dir: str | Path | None = None) -> Document:
    """把结构化描述渲染成 Document 对象"""
    style = _apply_template(DEFAULT_STYLE, spec.get("template"))
    base = Path(base_dir) if base_dir else None

    doc = Document()
    _setup_page(doc, style)

    if spec.get("cover"):
        _render_cover(doc, spec["cover"], style)
    if spec.get("toc"):
        _render_toc(doc)
    _render_header_footer(doc, spec)

    for block in spec.get("sections") or []:
        if not isinstance(block, dict):
            continue
        kind = str(block.get("type", "paragraph")).lower()
        if kind == "heading":
            _render_heading(doc, block, style)
        elif kind == "paragraph":
            _render_paragraph(doc, block, style)
        elif kind == "list":
            _render_list(doc, block, style)
        elif kind == "table":
            _render_table(doc, block, style)
        elif kind == "image":
            _render_image(doc, block, base)
        elif kind == "formula":
            _render_formula(doc, block)
        elif kind == "page_break":
            doc.add_page_break()
        elif kind == "section_break":
            _render_section_break(doc, block)
        else:
            logger.warning("未知的块类型，已跳过", type=kind)

    _enable_update_fields(doc)
    return doc


def render_to_file(spec: dict, out_path: str | Path, base_dir: str | Path | None = None) -> Path:
    """渲染并保存为 .docx"""
    target = Path(out_path)
    target.parent.mkdir(parents=True, exist_ok=True)
    doc = render(spec, base_dir)
    doc.save(str(target))
    return target
