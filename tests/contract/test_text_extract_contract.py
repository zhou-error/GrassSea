"""文本提取契约测试

背景：曾经用 ("utf-8","utf-8-sig","gbk","latin-1") 依次解码，
而 latin-1 对任意字节都不会失败，导致二进制文件被「成功」解成乱码
并以 completed 返回 —— 静默假成功。以下断言确保不会回退。
"""

from __future__ import annotations

import pytest

from shared.utils.text_extract import ExtractError, extract_text


def _make_xlsx(path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "型值表"
    ws.append(["站号", "半宽"])
    ws.append([0, 0])
    ws.append([1, 3.2])
    wb.save(path)


def _make_docx(path):
    from docx import Document

    doc = Document()
    doc.add_paragraph("设计说明书：垂线间长 120 m。")
    doc.save(path)


def test_text_file_ok(tmp_path):
    f = tmp_path / "a.txt"
    f.write_text("纯文本内容", encoding="utf-8")
    result = extract_text(f)
    assert result["format"] == "text"
    assert "纯文本内容" in result["text"]


def test_xlsx_parsed(tmp_path):
    f = tmp_path / "型值表.xlsx"
    _make_xlsx(f)
    result = extract_text(f)
    assert result["format"] == "xlsx"
    assert "型值表" in result["text"] and "3.2" in result["text"]


def test_docx_parsed(tmp_path):
    f = tmp_path / "说明书.docx"
    _make_docx(f)
    result = extract_text(f)
    assert result["format"] == "docx"
    assert "垂线间长" in result["text"]


def test_binary_rejected_not_mojibake(tmp_path):
    """核心断言：未知二进制必须报错，不能返回乱码且声称成功"""
    f = tmp_path / "blob.bin"
    f.write_bytes(bytes(range(256)) * 8)
    with pytest.raises(ExtractError):
        extract_text(f)


def test_legacy_doc_gives_actionable_hint(tmp_path):
    f = tmp_path / "旧文件.doc"
    f.write_bytes(b"\xd0\xcf\x11\xe0" + b"\x00" * 64)
    with pytest.raises(ExtractError) as exc:
        extract_text(f)
    assert "另存为" in str(exc.value)


def test_dwg_reports_not_implemented(tmp_path):
    f = tmp_path / "图纸.dwg"
    f.write_bytes(b"AC1032" + b"\x00" * 64)
    with pytest.raises(ExtractError) as exc:
        extract_text(f)
    assert "尚未实现" in str(exc.value)


def test_no_latin1_fallback_in_source():
    """防止有人把 latin-1 加回编码回退链"""
    from pathlib import Path

    source = Path("shared/utils/text_extract.py").read_text(encoding="utf-8")
    code = "\n".join(
        line for line in source.splitlines() if not line.strip().startswith("#")
    )
    assert '"latin-1"' not in code and "'latin-1'" not in code
